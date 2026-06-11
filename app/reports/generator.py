"""Report dataset builders + format renderers (CSV / Excel / PDF)."""
import csv
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from app.extensions import db
from app.models import (
    AssessmentAttempt,
    EmployeeSkill,
    Goal,
    LearningProgress,
    Skill,
    User,
)
from app.utils.scoring import compute_growth_score


def build_dataset(report_type, scope_ids, focus_employee=None):
    """Return (title, dataframe)."""
    if report_type == "employee" and focus_employee:
        return _employee_dataset(focus_employee)
    if report_type == "team":
        return _people_dataset("Team Performance Report", scope_ids)
    return _organization_dataset(scope_ids)


def _employee_dataset(employee_id):
    user = db.session.get(User, employee_id)
    score = compute_growth_score(employee_id, "monthly", persist=False)
    comp = score["components"]

    learning = LearningProgress.query.filter_by(employee_id=employee_id).all()
    completed = sum(1 for p in learning if p.status == "completed")
    attempts = AssessmentAttempt.query.filter(
        AssessmentAttempt.employee_id == employee_id,
        AssessmentAttempt.percentage.isnot(None),
    ).all()
    skills = (
        db.session.query(Skill.name, EmployeeSkill.level)
        .join(EmployeeSkill, EmployeeSkill.skill_id == Skill.id)
        .filter(EmployeeSkill.employee_id == employee_id)
        .all()
    )
    goals = Goal.query.filter_by(employee_id=employee_id).all()
    goals_done = sum(1 for g in goals if g.status == "completed")

    rows = [
        {"Metric": "Employee", "Value": user.full_name},
        {"Metric": "Email", "Value": user.email},
        {"Metric": "Modules assigned", "Value": len(learning)},
        {"Metric": "Modules completed", "Value": completed},
        {"Metric": "Assessments taken", "Value": len(attempts)},
        {"Metric": "Avg assessment %", "Value": comp["assessment"]},
        {"Metric": "Goals completed", "Value": f"{goals_done}/{len(goals)}"},
        {"Metric": "Skills tracked", "Value": len(skills)},
        {"Metric": "Learning score", "Value": comp["learning"]},
        {"Metric": "Goal score", "Value": comp["goal"]},
        {"Metric": "Skill growth score", "Value": comp["skill"]},
        {"Metric": "Consistency score", "Value": comp["consistency"]},
        {"Metric": "Overall Growth Score", "Value": score["overall"]},
        {"Metric": "Category", "Value": score["category"].replace("_", " ").title()},
    ]
    return f"Employee Report — {user.full_name}", pd.DataFrame(rows)


def _people_dataset(title, scope_ids):
    employees = User.query.filter(
        User.id.in_(scope_ids), User.role == "employee"
    ).order_by(User.full_name).all()
    rows = []
    for emp in employees:
        score = compute_growth_score(emp.id, "monthly", persist=False)
        comp = score["components"]
        rows.append(
            {
                "Employee": emp.full_name,
                "Email": emp.email,
                "Learning": comp["learning"],
                "Assessment": comp["assessment"],
                "Goals": comp["goal"],
                "Skill": comp["skill"],
                "Consistency": comp["consistency"],
                "Overall": score["overall"],
                "Category": score["category"].replace("_", " ").title(),
            }
        )
    return title, pd.DataFrame(rows)


def _organization_dataset(scope_ids):
    title, df = _people_dataset("Organization Workforce Report", scope_ids)
    if df.empty:
        return title, df
    summary = pd.DataFrame(
        [
            {"Employee": "— ORG AVERAGE —", "Email": "",
             "Learning": round(df["Learning"].mean(), 1),
             "Assessment": round(df["Assessment"].mean(), 1),
             "Goals": round(df["Goals"].mean(), 1),
             "Skill": round(df["Skill"].mean(), 1),
             "Consistency": round(df["Consistency"].mean(), 1),
             "Overall": round(df["Overall"].mean(), 1),
             "Category": ""}
        ]
    )
    return title, pd.concat([df, summary], ignore_index=True)


# --------------------------------------------------------------------------- #
# Renderers
# --------------------------------------------------------------------------- #
def to_csv(df):
    buf = io.StringIO()
    df.to_csv(buf, index=False, quoting=csv.QUOTE_MINIMAL)
    return io.BytesIO(buf.getvalue().encode("utf-8"))


def to_excel(title, df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([title])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([])
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True)):
        ws.append(row)
        if r_idx == 0:
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 4, 45)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def to_pdf(title, df):
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), title=title)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    data = [list(df.columns)] + df.astype(str).values.tolist()
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF2F7")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    out.seek(0)
    return out


def render(report_type, fmt, scope_ids, focus_employee=None):
    title, df = build_dataset(report_type, scope_ids, focus_employee)
    if fmt == "csv":
        return to_csv(df), "text/csv", "csv"
    if fmt == "excel":
        return (
            to_excel(title, df),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    return to_pdf(title, df), "application/pdf", "pdf"
